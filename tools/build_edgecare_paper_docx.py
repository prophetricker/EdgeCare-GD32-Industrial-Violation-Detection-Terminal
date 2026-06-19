#!/usr/bin/env python3
"""Build the EdgeCare technical paper DOCX from the Markdown draft."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


TITLE = "EdgeCare：基于GD32H759的端侧工业危险区闯入检测终端"
CONTEST_HEADER = "第二十一届中国研究生电子设计竞赛"
MD_PATH = Path("doc/competition_materials/EdgeCare_技术论文_初稿.md")
DOCX_PATH = Path("doc/competition_materials/EdgeCare_技术论文_初稿.docx")


def set_run_font(run, size: int | float, bold: bool = False, font: str = "宋体") -> None:
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold


def set_paragraph_format(paragraph, size: int | float = 12, bold: bool = False, align=None) -> None:
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.line_spacing = Pt(20)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    for run in paragraph.runs:
        set_run_font(run, size, bold)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def restart_page_number(section, start: int = 1) -> None:
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), str(start))


def set_header_footer(section, include_page_number: bool) -> None:
    header = section.header
    header.is_linked_to_previous = False
    if header.paragraphs:
        paragraph = header.paragraphs[0]
    else:
        paragraph = header.add_paragraph()
    paragraph.text = CONTEST_HEADER + "\n" + TITLE
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in paragraph.runs:
        set_run_font(run, 10.5)

    footer = section.footer
    footer.is_linked_to_previous = False
    if include_page_number:
        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        paragraph.text = ""
        add_page_number(paragraph)
        for run in paragraph.runs:
            set_run_font(run, 10.5)


def set_section_layout(section, include_page_number: bool) -> None:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)
    set_header_footer(section, include_page_number)


def add_centered_paragraph(doc: Document, text: str, size: int | float = 12, bold: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_run_font(run, size, bold, "黑体" if bold else "宋体")
    paragraph.paragraph_format.line_spacing = Pt(20)


def parse_markdown_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        row = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in row):
            rows.append(row)
        i += 1
    return rows, i


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text = row[c_idx] if c_idx < len(row) else ""
            cell.text = text
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                set_paragraph_format(paragraph, 10.5, bold=(r_idx == 0))
            if r_idx == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "D9D9D9")
                cell._tc.get_or_add_tcPr().append(shading)


def add_heading(doc: Document, text: str, level: int) -> None:
    if level == 1:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        set_run_font(run, 18, True, "黑体")
    elif level == 2:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(text)
        set_run_font(run, 14, True, "黑体")
    else:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(text)
        set_run_font(run, 12, True, "黑体")
    paragraph.paragraph_format.line_spacing = Pt(20)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)


def add_body_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    run = paragraph.add_run(text)
    set_run_font(run, 12)
    paragraph.paragraph_format.line_spacing = Pt(20)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)


def add_code_block(doc: Document, block: list[str]) -> None:
    for line in block:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.5)
        run = paragraph.add_run(line)
        set_run_font(run, 9, False, "Consolas")
        paragraph.paragraph_format.line_spacing = Pt(12)


def build_document(md_text: str) -> Document:
    doc = Document()
    set_section_layout(doc.sections[0], include_page_number=False)

    lines = md_text.splitlines()
    body_start = lines.index("## 封面") if "## 封面" in lines else 0
    del body_start

    add_centered_paragraph(doc, CONTEST_HEADER, 18, True)
    add_centered_paragraph(doc, "技术论文", 26, True)
    for _ in range(2):
        doc.add_paragraph()
    cover_lines = [
        "论文题目：",
        "中文：EdgeCare：基于GD32H759的端侧工业危险区闯入检测终端",
        "英文：EdgeCare: An Edge Industrial Danger-Zone Intrusion Detection Terminal Based on GD32H759",
        "参赛单位：【待填写】",
        "队伍名称：【待填写】",
        "指导老师：【待填写】",
        "参赛队员：【待填写】",
        "完成时间：【待填写】",
    ]
    for line in cover_lines:
        add_centered_paragraph(doc, line, 15, True)

    doc.add_section(WD_SECTION_START.NEW_PAGE)
    set_section_layout(doc.sections[-1], include_page_number=False)

    toc_heading = doc.add_paragraph()
    toc_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = toc_heading.add_run("目    录")
    set_run_font(run, 16, True, "黑体")
    for line in lines:
        if line.startswith("第") or re.match(r"^\d+\.\d+", line):
            p = doc.add_paragraph(line)
            set_paragraph_format(p, 12)

    doc.add_section(WD_SECTION_START.NEW_PAGE)
    set_section_layout(doc.sections[-1], include_page_number=False)

    content_started = False
    page_number_started = False
    i = 0
    in_code = False
    code_block: list[str] = []
    skip_sections = {"# EdgeCare：基于GD32H759的端侧工业危险区闯入检测终端", "## 封面", "## 目录"}

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if line in skip_sections:
            i += 1
            continue
        if line.startswith("```"):
            if in_code:
                add_code_block(doc, code_block)
                code_block = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_block.append(raw)
            i += 1
            continue
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            rows, next_i = parse_markdown_table(lines, i)
            add_table(doc, rows)
            i = next_i
            continue
        if line.startswith("## "):
            title = line[3:]
            if title in {"中文摘要", "English Abstract", "参考文献"}:
                if content_started:
                    doc.add_section(WD_SECTION_START.NEW_PAGE)
                    set_section_layout(doc.sections[-1], include_page_number=(title == "参考文献"))
                content_started = True
                add_heading(doc, title, 1)
            elif title.startswith("第"):
                if content_started:
                    doc.add_section(WD_SECTION_START.NEW_PAGE)
                    set_section_layout(doc.sections[-1], include_page_number=True)
                    if not page_number_started:
                        restart_page_number(doc.sections[-1], 1)
                        page_number_started = True
                content_started = True
                add_heading(doc, title, 1)
            else:
                add_heading(doc, title, 1)
            i += 1
            continue
        if line.startswith("### "):
            add_heading(doc, line[4:], 2)
            i += 1
            continue
        if line.startswith("#### "):
            add_heading(doc, line[5:], 3)
            i += 1
            continue
        if re.match(r"^(图|表)\d+-\d+ ", line):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line)
            set_run_font(run, 10.5)
            i += 1
            continue
        if line.startswith("资料来源："):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line)
            set_run_font(run, 10.5)
            i += 1
            continue
        if line.startswith("[") and line[1:2].isdigit():
            p = doc.add_paragraph(line)
            set_paragraph_format(p, 10.5)
            i += 1
            continue
        add_body_paragraph(doc, line)
        i += 1

    return doc


def main() -> int:
    if not MD_PATH.exists():
        print(f"missing markdown: {MD_PATH}", file=sys.stderr)
        return 1
    doc = build_document(MD_PATH.read_text(encoding="utf-8"))
    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
